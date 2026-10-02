import os
import time
import pdfplumber
from docx import Document
import openpyxl
import chromadb
from openai import OpenAI
from tqdm import tqdm

# ---------- CONFIG ----------
BASE_DIR      = os.path.dirname(os.path.abspath(__file__))
KEY_FILE      = os.path.join(BASE_DIR, "openai_key.txt")
DOCS_DIR      = os.path.join(BASE_DIR, "test_docs")
DB_DIR        = os.path.join(BASE_DIR, "chroma_db")
COLLECTION    = "gkhep_docs"

EMBED_MODEL   = "text-embedding-3-small"
CHUNK_SIZE    = 1000
CHUNK_OVERLAP = 200
BATCH_SIZE    = 100

# ---------- SETUP ----------
api_key = open(KEY_FILE, "r", encoding="utf-8").read().strip()
client  = OpenAI(api_key=api_key)

chroma_client = chromadb.PersistentClient(path=DB_DIR)
collection = chroma_client.get_or_create_collection(name=COLLECTION)

# ---------- FILE READERS ----------
def read_pdf(path):
    out = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            t = page.extract_text() or ""
            if t.strip():
                out.append(f"[Page {i}]\n{t}")
    return "\n\n".join(out)

def read_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for tbl in doc.tables:
        for row in tbl.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)

def read_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        out.append(f"=== Sheet: {sheet} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                out.append(" | ".join(cells))
    return "\n".join(out)

def read_any(path):
    ext = path.lower()
    if ext.endswith(".pdf"):
        return read_pdf(path)
    if ext.endswith(".docx"):
        return read_docx(path)
    if ext.endswith((".xlsx", ".xlsm")):
        return read_xlsx(path)
    return None

# ---------- CHUNKING ----------
def chunk_text(text, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    chunks = []
    start = 0
    while start < len(text):
        end = start + size
        chunks.append(text[start:end])
        start = end - overlap
    return [c for c in chunks if c.strip()]

# ---------- COLLECT FILES ----------
files = []
for root, _, names in os.walk(DOCS_DIR):
    for n in names:
        if n.lower().endswith((".pdf", ".docx", ".xlsx", ".xlsm")):
            files.append(os.path.join(root, n))

print(f"\n📁 Found {len(files)} document(s) in {DOCS_DIR}\n")
if not files:
    print("❌ No documents found. Put some PDF/DOCX/XLSX files into test_docs/")
    raise SystemExit

# ---------- PROCESS ----------
all_ids, all_texts, all_meta = [], [], []

for path in tqdm(files, desc="Reading files"):
    fname = os.path.basename(path)
    try:
        text = read_any(path)
        if not text or not text.strip():
            print(f"  ⚠️  Empty: {fname}")
            continue
        pieces = chunk_text(text)
        for i, piece in enumerate(pieces):
            uid = f"{fname}__chunk_{i}"
            all_ids.append(uid)
            all_texts.append(piece)
            all_meta.append({"source": fname, "chunk": i})
        print(f"  ✓ {fname}: {len(pieces)} chunks")
    except Exception as e:
        print(f"  ❌ {fname}: {e}")

print(f"\n🧮 Total chunks: {len(all_texts)}")

# ---------- EMBEDDINGS + STORE ----------
print("\n🔢 Creating embeddings and storing in ChromaDB...\n")

for i in range(0, len(all_texts), BATCH_SIZE):
    batch_texts = all_texts[i:i+BATCH_SIZE]
    batch_ids   = all_ids[i:i+BATCH_SIZE]
    batch_meta  = all_meta[i:i+BATCH_SIZE]

    resp = client.embeddings.create(
        model=EMBED_MODEL,
        input=batch_texts,
    )
    batch_embs = [d.embedding for d in resp.data]

    collection.upsert(
        ids=batch_ids,
        documents=batch_texts,
        metadatas=batch_meta,
        embeddings=batch_embs,
    )
    print(f"  ✓ Stored batch {i//BATCH_SIZE + 1} ({len(batch_texts)} chunks)")

print(f"\n✅ Indexing complete.")
print(f"   Database: {DB_DIR}")
print(f"   Total items in collection: {collection.count()}")