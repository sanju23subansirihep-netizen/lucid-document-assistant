import os
import re
import json
import streamlit as st
import chromadb
import pdfplumber
from docx import Document
import openpyxl
from openai import OpenAI
from datetime import datetime

# ---------- CONFIG ----------
BASE_DIR    = os.path.dirname(os.path.abspath(__file__))
KEY_FILE    = os.path.join(BASE_DIR, "openai_key.txt")
DOCS_DIR    = os.path.join(BASE_DIR, "test_docs")
DB_DIR      = os.path.join(BASE_DIR, "chroma_db")
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
CHATS_FILE  = os.path.join(BASE_DIR, "chat_history.json")
TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
os.makedirs(UPLOADS_DIR, exist_ok=True)

EMBED_MODEL   = "text-embedding-3-small"
CHAT_MODEL    = "gpt-4o-mini"
CHUNK_SIZE    = 1000
CHUNK_OVERLAP = 200
BATCH_SIZE    = 100
TOP_K         = 5

# ---------- PAGE ----------
st.set_page_config(
    page_title="Lucid",
    page_icon="◆",
    layout="wide",
)

# ---------- CHAT HISTORY ----------
def load_chat_sessions():
    if os.path.exists(CHATS_FILE):
        try:
            with open(CHATS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_chat_sessions(sessions):
    try:
        with open(CHATS_FILE, "w", encoding="utf-8") as f:
            json.dump(sessions, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

def new_session_id():
    return datetime.now().strftime("%Y%m%d_%H%M%S")

def make_title(messages):
    first_user = next(
        (m["content"] for m in messages if m["role"] == "user"),
        "Untitled chat"
    )
    return first_user[:40] + ("..." if len(first_user) > 40 else "")

def save_current_chat():
    if not st.session_state.messages:
        return
    st.session_state.chat_sessions[st.session_state.current_session_id] = {
        "title": make_title(st.session_state.messages),
        "messages": st.session_state.messages,
        "created": st.session_state.current_session_id,
    }
    save_chat_sessions(st.session_state.chat_sessions)

# ---------- LOAD CSS ----------
def load_css():
    css_path = os.path.join(BASE_DIR, "style.css")
    if os.path.exists(css_path):
        with open(css_path, "r", encoding="utf-8") as f:
            st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)

load_css()

# ---------- INIT ----------
@st.cache_resource
def get_openai():
    api_key = open(KEY_FILE, "r", encoding="utf-8").read().strip()
    return OpenAI(api_key=api_key)

@st.cache_resource
def get_chroma():
    return chromadb.PersistentClient(path=DB_DIR)

try:
    oai = get_openai()
    chroma = get_chroma()
except Exception as e:
    st.error(f"Startup error: {e}")
    st.stop()

# ---------- SESSION STATE ----------
if "dark_mode" not in st.session_state:
    st.session_state.dark_mode = False

if "chat_sessions" not in st.session_state:
    st.session_state.chat_sessions = load_chat_sessions()

if "current_session_id" not in st.session_state:
    st.session_state.current_session_id = new_session_id()

if "messages" not in st.session_state:
    st.session_state.messages = []

if "collection_name" not in st.session_state:
    st.session_state.collection_name = "gkhep_docs"

# ---------- COLLECTIONS ----------
def list_collections():
    cols = chroma.list_collections()
    return [c.name for c in cols] if cols else ["gkhep_docs"]

def get_collection(name=None):
    if name is None:
        name = st.session_state.collection_name
    return chroma.get_or_create_collection(name=name)

# ---------- FILE READERS ----------
def read_pdf_ocr(path):
    import pytesseract
    import pypdfium2 as pdfium

    if os.path.exists(TESSERACT_PATH):
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH

    pdf = pdfium.PdfDocument(path)
    out = []
    for i in range(len(pdf)):
        page = pdf[i]
        bitmap = page.render(scale=200 / 72)
        pil_image = bitmap.to_pil()
        text = pytesseract.image_to_string(pil_image)
        if text.strip():
            out.append(f"[Page {i+1} - OCR]\n{text}")
    pdf.close()
    return "\n\n".join(out)

def read_pdf(path):
    out = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            t = page.extract_text() or ""
            if t.strip():
                out.append(f"[Page {i}]\n{t}")
    text = "\n\n".join(out)
    if len(text.strip()) < 50:
        try:
            ocr_text = read_pdf_ocr(path)
            if ocr_text.strip():
                return ocr_text
        except Exception as e:
            print(f"[OCR failed] {path}: {e}")
    return text

def read_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for tbl in doc.tables:
        for row in tbl.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)

def read_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        out.append(f"=== Sheet: {sheet} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                out.append(" | ".join(cells))
    return "\n".join(out)

def read_any(path):
    ext = path.lower()
    if ext.endswith(".pdf"):
        return read_pdf(path)
    if ext.endswith(".docx"):
        return read_docx(path)
    if ext.endswith((".xlsx", ".xlsm")):
        return read_xlsx(path)
    return None

def chunk_text(text, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    chunks, start = [], 0
    while start < len(text):
        end = start + size
        chunks.append(text[start:end])
        start = end - overlap
    return [c for c in chunks if c.strip()]

def safe_filename(name):
    return re.sub(r'[&\\/:*?"<>|]', "_", name).strip()

# ---------- INDEXING ----------
def index_file(path, display_name, collection):
    text = read_any(path)
    if not text or not text.strip():
        return 0, "Empty file (may be a scanned PDF)"
    pieces = chunk_text(text)
    if not pieces:
        return 0, "No text extracted"
    ids, texts, metas = [], [], []
    for i, piece in enumerate(pieces):
        ids.append(f"{display_name}__chunk_{i}")
        texts.append(piece)
        metas.append({"source": display_name, "chunk": i})
    total = 0
    for j in range(0, len(texts), BATCH_SIZE):
        bt = texts[j:j+BATCH_SIZE]
        bi = ids[j:j+BATCH_SIZE]
        bm = metas[j:j+BATCH_SIZE]
        resp = oai.embeddings.create(model=EMBED_MODEL, input=bt)
        embs = [d.embedding for d in resp.data]
        collection.upsert(ids=bi, documents=bt, metadatas=bm, embeddings=embs)
        total += len(bt)
    return total, "OK"

# ---------- Q&A ----------
def find_chunks(collection, question, k=TOP_K):
    q_emb = oai.embeddings.create(
        model=EMBED_MODEL, input=[question]
    ).data[0].embedding
    results = collection.query(query_embeddings=[q_emb], n_results=k)
    return list(zip(
        results["documents"][0],
        results["metadatas"][0],
        results["distances"][0],
    ))

def build_prompts(hits, question):
    context = "\n\n---\n\n".join(
        f"[From: {m['source']}, chunk {m['chunk']}]\n{d}" for d, m, _ in hits
    )
    sys_prompt = (
        "You are Lucid, an assistant for project documents. "
        "Answer ONLY using the context below. If not found, say "
        "'I could not find this in the provided documents.' "
        "Cite source file names at the end."
    )
    user_prompt = f"CONTEXT:\n{context}\n\nQUESTION:\n{question}\n\nANSWER:"
    return sys_prompt, user_prompt

# ---------- PDF EXPORT ----------
def export_chat_pdf(messages):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Lucid - Chat Export", ln=True)
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 6, f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M')}", ln=True)
    pdf.ln(3)
    for m in messages:
        pdf.set_font("Helvetica", "B", 11)
        role = "You" if m["role"] == "user" else "Lucid"
        pdf.cell(0, 7, f"{role}:", ln=True)
        pdf.set_font("Helvetica", "", 10)
        text = m["content"].encode("latin-1", "replace").decode("latin-1")
        pdf.multi_cell(0, 5, text)
        pdf.ln(2)
    return bytes(pdf.output())

# ---------- SIDEBAR ----------
with st.sidebar:
    st.markdown("""
    <div class="lucid-brand">
        <div class="monogram">L</div>
        <div>
            <div class="brand-name">Lucid</div>
            <span class="brand-tag">Clarity from complexity</span>
        </div>
    </div>
    """, unsafe_allow_html=True)

    dark = st.toggle("Dark mode", value=st.session_state.dark_mode)
    if dark != st.session_state.dark_mode:
        st.session_state.dark_mode = dark
        st.rerun()

    if st.button("New chat", use_container_width=True):
        save_current_chat()
        st.session_state.current_session_id = new_session_id()
        st.session_state.messages = []
        st.rerun()

    if st.session_state.chat_sessions:
        with st.expander(f"Previous chats ({len(st.session_state.chat_sessions)})"):
            sorted_sessions = sorted(
                st.session_state.chat_sessions.items(),
                key=lambda x: x[0],
                reverse=True,
            )
            for sid, sess in sorted_sessions[:20]:
                col1, col2 = st.columns([5, 1])
                label = f"{sess.get('title', 'Untitled')}"
                if col1.button(label, key=f"load_{sid}", use_container_width=True):
                    save_current_chat()
                    st.session_state.messages = sess["messages"]
                    st.session_state.current_session_id = sid
                    st.rerun()
                if col2.button("X", key=f"del_chat_{sid}"):
                    del st.session_state.chat_sessions[sid]
                    save_chat_sessions(st.session_state.chat_sessions)
                    st.rerun()

    st.divider()

    st.markdown("### Workspace")
    cols = list_collections()
    if st.session_state.collection_name not in cols:
        cols.append(st.session_state.collection_name)
    selected = st.selectbox(
        "Active collection",
        options=cols,
        index=cols.index(st.session_state.collection_name),
    )
    if selected != st.session_state.collection_name:
        st.session_state.collection_name = selected
        st.rerun()

    new_col = st.text_input("Create new collection", placeholder="e.g. contracts")
    if st.button("Create", use_container_width=True) and new_col.strip():
        chroma.get_or_create_collection(name=new_col.strip())
        st.session_state.collection_name = new_col.strip()
        st.success(f"Created '{new_col}'")
        st.rerun()

    collection = get_collection()

    st.divider()
    st.markdown("### Upload documents")

    uploaded = st.file_uploader(
        "Drop files here",
        type=["pdf", "docx", "xlsx", "xlsm"],
        accept_multiple_files=True,
        label_visibility="collapsed",
    )

    if uploaded:
        if st.button(f"Index {len(uploaded)} file(s)", use_container_width=True):
            progress = st.progress(0.0)
            total_chunks = 0
            errors = []
            for i, uf in enumerate(uploaded):
                clean = safe_filename(uf.name)
                save_path = os.path.join(UPLOADS_DIR, clean)
                try:
                    with open(save_path, "wb") as f:
                        f.write(uf.getbuffer())
                except Exception as e:
                    errors.append(f"Save failed for {uf.name}: {e}")
                    progress.progress((i + 1) / len(uploaded))
                    continue
                try:
                    chunks, status = index_file(save_path, clean, collection)
                    if status == "OK":
                        total_chunks += chunks
                    else:
                        errors.append(f"{clean}: {status}")
                except Exception as e:
                    errors.append(f"{clean}: {e}")
                progress.progress((i + 1) / len(uploaded))

            if total_chunks > 0:
                st.success(f"Indexed {total_chunks} chunks.")
            for err in errors:
                st.error(err)
            if total_chunks > 0:
                st.rerun()

    st.divider()
    st.markdown("### Indexed documents")

    try:
        all_meta = collection.get(include=["metadatas"])["metadatas"]
        sources = sorted(set(m["source"] for m in all_meta))
    except Exception:
        sources = []

    if not sources:
        st.caption("No documents indexed yet.")
    else:
        for src in sources:
            col1, col2 = st.columns([4, 1])
            col1.markdown(f"· {src}")
            if col2.button("X", key=f"del_{src}"):
                all_items = collection.get(include=["metadatas"])
                ids_to_del = [
                    all_items["ids"][i]
                    for i, m in enumerate(all_items["metadatas"])
                    if m["source"] == src
                ]
                if ids_to_del:
                    collection.delete(ids=ids_to_del)
                st.success(f"Deleted {src}")
                st.rerun()

    st.divider()

    try:
        _all_meta = collection.get(include=["metadatas"])["metadatas"]
        _source_count = len(set(m["source"] for m in _all_meta))
        _chunk_count = len(_all_meta)
    except Exception:
        _source_count = 0
        _chunk_count = 0
    st.caption(f"{_source_count} documents · {_chunk_count} chunks")

    if st.button("Re-index test_docs folder", use_container_width=True):
        with st.spinner("Indexing test_docs..."):
            count = 0
            for root, _, names in os.walk(DOCS_DIR):
                for n in names:
                    if n.lower().endswith((".pdf", ".docx", ".xlsx", ".xlsm")):
                        p = os.path.join(root, n)
                        try:
                            chunks, status = index_file(p, n, collection)
                            count += chunks
                        except Exception as e:
                            st.warning(f"{n}: {e}")
            st.success(f"Indexed {count} chunks.")
            st.rerun()

# ---------- DARK MODE OVERRIDE ----------
if st.session_state.get("dark_mode", False):
    st.markdown("""
<style>
.main, [data-testid="stAppViewContainer"], .block-container {
    background: #0B1524 !important;
}
h1, h2, h3, p, span, label, .stMarkdown {
    color: #E8EEF6 !important;
}
.lucid-hero .wordmark {
    color: #FFFFFF !important;
    -webkit-text-fill-color: #FFFFFF !important;
}
.lucid-hero .tagline {
    color: #8BA3BE !important;
}
[data-testid="stChatMessage"] {
    background: #12203A !important;
    border-color: #1F3456 !important;
}
[data-testid="stChatMessage"] p {
    color: #E8EEF6 !important;
}
[data-testid="stChatInput"] {
    background: #12203A !important;
    border-color: #1F3456 !important;
}
[data-testid="stChatInput"] textarea {
    color: #E8EEF6 !important;
}
.main .stButton > button {
    background: #12203A !important;
    color: #E8EEF6 !important;
    border-color: #1F3456 !important;
}
.streamlit-expanderHeader, details summary {
    background: #12203A !important;
    color: #E8EEF6 !important;
    border-color: #1F3456 !important;
}
</style>
""", unsafe_allow_html=True)

# ---------- MAIN HERO ----------
st.markdown("""
<div class="lucid-hero">
    <div class="wordmark">Lucid</div>
    <span class="tagline">Clarity from complexity</span>
</div>
""", unsafe_allow_html=True)

# ---------- EXPORT BUTTON ----------
if st.session_state.messages:
    col1, col2 = st.columns([6, 1])
    with col2:
        pdf_bytes = export_chat_pdf(st.session_state.messages)
        st.download_button(
            "Export PDF",
            data=pdf_bytes,
            file_name=f"lucid_chat_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
            mime="application/pdf",
        )

# ---------- CHAT MESSAGES ----------
AVATARS = {"user": "👤", "assistant": "◆"}

for m in st.session_state.messages:
    with st.chat_message(m["role"], avatar=AVATARS.get(m["role"], "◆")):
        st.markdown(m["content"])
        if m.get("sources"):
            with st.expander("Sources used"):
                for s in m["sources"]:
                    st.write(f"· **{s['source']}** — chunk {s['chunk']} (score {s['score']:.3f})")

# ---------- CHAT INPUT ----------
prompt = st.chat_input("Ask anything of your documents...")

if prompt:
    st.session_state.messages.append({"role": "user", "content": prompt})
    with st.chat_message("user", avatar="👤"):
        st.markdown(prompt)

    with st.chat_message("assistant", avatar="◆"):
        try:
            hits = find_chunks(collection, prompt)
            sys_prompt, user_prompt = build_prompts(hits, prompt)

            stream = oai.chat.completions.create(
                model=CHAT_MODEL,
                messages=[
                    {"role": "system", "content": sys_prompt},
                    {"role": "user",   "content": user_prompt},
                ],
                temperature=0.2,
                stream=True,
            )

            full_answer = st.write_stream(
                (chunk.choices[0].delta.content or "")
                for chunk in stream
                if chunk.choices[0].delta.content
            )

            sources_list = [
                {"source": m["source"], "chunk": m["chunk"], "score": d}
                for _, m, d in hits
            ]
            with st.expander("Sources used"):
                for s in sources_list:
                    st.write(
                        f"· **{s['source']}** — chunk {s['chunk']} "
                        f"(score {s['score']:.3f})"
                    )
            st.session_state.messages.append({
                "role": "assistant",
                "content": full_answer,
                "sources": sources_list,
            })
            save_current_chat()
        except Exception as e:
            st.error(f"Error: {e}")